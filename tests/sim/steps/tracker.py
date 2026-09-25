"""Tracker v2: ENT-001..076 (сим 1, регрессийн төлөвтэй) + SIM2-NNN (сим 2.0). python3 steps/tracker.py → /mnt/user-data/outputs/entry_sim2_tracker.xlsx"""
import json, os, sys, subprocess
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.environ.get('SIM_TRACKER_SRC', 'entry_simulation_issues.xlsx')
OUT = os.environ.get('SIM_TRACKER_OUT', 'entry_sim2_tracker.xlsx')
REG = {  # ENT id → (төлөв, нотолгоо)
    'ENT-069': ('Засагдсан ✓', 'get_counterparty_balance B/C 2026-06/2025-12 амжилттай (aging-тай)'),
    'ENT-018': ('Засагдсан ✓', 'C: PO-гүй АП бараатай мөр (Дотоод Нийлүүлэгч) сар бүр автоматаар үнэлэгдэв — блоклогдсон 0'),
    'ENT-023': ('Засагдсан ✓', 'C: USD/CNY/EUR сар бүрийн FX тэгшитгэл oracle-тэй таарав (2025-01..08)'),
    'ENT-011': ('Засагдсан ✓', 'C: Голомт USD нээлт 38,500 USD → GL 131,687,710₮ (ханштай)'),
    'ENT-012': ('Засагдсан ✓', 'C: нээлтийн журнал 2024-12-31 огноотой'),
    'ENT-024': ('Засагдсан ✓', 'C: НӨАТ тайлангийн гаралт/оролт/шилжүүлэх дүн oracle-тэй таарав (SIM2-027 label-ын асуудал үлдсэн)'),
    'ENT-035': ('Засагдсан ✓', 'C: VAT-25-0000NN тооцооны журнал сарын сүүлийн өдрөөр'),
    'ENT-072': ('Засагдсан ✓', 'B/C get_balance_sheet «Тэнцэл ✓»'),
    'ENT-047': ('Хэсэгчлэн — SIM2-043', 'Касс 10000001 мөнгөнд орсон; гэвч АР төлбөр «Авлагын өөрчлөлт» хэвээр'),
    'ENT-067': ('Засагдсан ✓', 'C: вэб PO хаалт «2025-01-31 өдрөөр хаах уу?» — сонгосон үеийн огноо'),
    'ENT-071': ('Засагдсан ✓', 'C: USD PO-той гаалийн нэхэмжлэх 511,000 MNT'),
    'ENT-066': ('Засагдсан ✓', 'B: Соната хасалт нээлтийн + системийн хур. элэгдэл 15,750,000 бүхэлдээ хаагдав'),
    'ENT-002': ('Засагдсан ✓', 'Вэб картанд «Нээлтийн хуримтлагдсан элэгдэл / cut-off»; B тавиур хугацаа дуусаад зогссон (19 идэвхтэй, 18 элэгдсэн)'),
    'ENT-043': ('Засагдсан ✓', 'C 2025-10: өртөггүй гар орлого IMP-021 → «шинээр үнэлэгдсэн 91, БЛОКЛОГДСОН (1)» — бусад бараа үнэлэгдэв; /costing-оос өртөг бөглөөд дахин тооцов'),
    'ENT-022': ('Засагдсан ✓', 'C/D: борлуулалтын COGS 61100000'),
    'ENT-044': ('Засагдсан ✓', 'C: WH-01→WH-02/03 шилжүүлэг дунджаар үнэлэгдэж дараагийн зарлага/тооллого блоклогдоогүй'),
    'ENT-029': ('Хойшлогдсон', 'Credit note — docs/product/2026-09-audit-followup-proposal.md'),
    'ENT-065': ('Хойшлогдсон', 'ECL — B-д гар журналаар (SIM2-038 анхааруулгагүй)'),
    'ENT-003': ('Хойшлогдсон', 'Нээлтийн барааны өртөг — C-д /costing гараар (SIM2-007)'),
    'ENT-064': ('Хойшлогдсон', 'Short close'),
    'ENT-027': ('Засагдсан ✓', 'C: 2025-11-31 → «Огноо хуанлид байхгүй»'),
    'ENT-013': ('Засагдсан ✓', 'create_journal_voucher-д currency, exchangeRate талбар бий'),
    'ENT-046': ('Засагдсан ✓', 'C: нээлтийн журнал 20000001/21000001 мөртэй ч давхар ноорог карт үүсээгүй (12 карт)'),
    'ENT-070': ('Засагдсан ✓', 'C mcplog: SQL/UUID алдаа 0 (JS runtime алдаа ил гарах нь үлдсэн — SIM2-016)'),
    'ENT-048': ('Хэсэгчлэн', 'Мөрийн нэр зөв болсон ч 87100004 тооллогын дутагдал «Санхүүгийн зардал» бүлэгт хэвээр'),
    'ENT-020': ('Хэсэгчлэн — SIM2-006', 'FC дансыг ₮-өөр харьцуулахаа больж «ТОДОРХОЙГҮЙ — нээлтийн ханш алга» гэнэ; тулгалт хийгдэхгүй хэвээр'),
    'ENT-052': ('Засагдсан ✓', 'C: шилжсэн кредит тайланд орж, тооцоо out-ыг л хаана'),
    'ENT-039': ('Засагдсан ✓', 'D вэб: «Мөр #2: бараатай мөрөнд агуулах заавал сонгоно»'),
    'ENT-040': ('Хэсэгчлэн', 'D вэб: барааны нүд хайлттай болсон ч «BM-012» гэж бичихэд BM-002 сонгогдов (кодоор нарийн тааруулахгүй); нэрээр зөв'),
    'ENT-021': ('Хэсэгчлэн', 'C жилийн хаалтын дараа reconcile: БАРАА МАТЕРИАЛ OK (худал зөрүүгүй) боловч КЛИРИНГ хэсэг нээлтийн 14000099 мөрийг (1.26B) «ХААГДААГҮЙ» гэсээр — мөр тус бүрээр таарууллаа гэж үзэхгүй'),
    'ENT-073': ('Засагдсан ✓', 'B/C жилийн хаалтын дараа reconcile_modules COGS/тооллого/14000099-д худал зөрүү гаргаагүй'),
    'ENT-001': ('Засагдсан ✓', 'B/C биет ҮХ 20000001/20000002 default'),
    'ENT-049': ('Засагдсан ✓', 'B ҮХ тайланд нээлтийн хур. элэгдэл харагдана (fa detail 15,750,000)'),
}
wb = openpyxl.load_workbook(SRC)
ws_ent = wb['Issues']
hdr = [c.value for c in ws_ent[1]]
ent = [dict(zip(hdr, r)) for r in ws_ent.iter_rows(min_row=2, values_only=True) if r[0]]
sim2 = [json.loads(l) for l in open(os.path.join(ROOT, 'out', 'issues2.jsonl'))]

out = openpyxl.Workbook()
F = 'Arial'
def style_header(ws, n):
    for c in ws[1][:n]:
        c.font = Font(name=F, bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F3864'); c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'
SEVFILL = {'S1': 'F8CBAD', 'S2': 'FFE699', 'S3': 'DDEBF7', 'S4': 'E2EFDA'}

# ── Summary ──
s = out.active; s.title = 'Summary'
s['A1'] = 'Entry — симуляцийн tracker v2 (сим 1 + сим 2.0)'; s['A1'].font = Font(name=F, bold=True, size=14)
s['A2'] = 'Хоёр хуудас: ENT (сим 1, 76 асуудал, регрессийн төлөвтэй) ба SIM2 (сим 2.0 шинэ олдвор). Тоо нь COUNTIF томьёогоор — мөр нэмэхэд шинэчлэгдэнэ.'; s['A2'].font = Font(name=F, italic=True)
s.append([]); s.append(['Зэрэг', 'ENT (сим 1)', 'SIM2 (сим 2.0)', 'Нийт']);
for c in s[4]: c.font = Font(name=F, bold=True)
for i, sev in enumerate(['S1', 'S2', 'S3', 'S4']):
    r = 5 + i
    s.cell(r, 1, sev).font = Font(name=F)
    s.cell(r, 2, f'=COUNTIF(ENT!$B:$B,A{r})'); s.cell(r, 3, f'=COUNTIF(SIM2!$B:$B,A{r})'); s.cell(r, 4, f'=B{r}+C{r}')
s.cell(9, 1, 'Нийт').font = Font(name=F, bold=True); s.cell(9, 2, '=SUM(B5:B8)'); s.cell(9, 3, '=SUM(C5:C8)'); s.cell(9, 4, '=SUM(D5:D8)')
s.append([]); s.append(['ENT регрессийн төлөв (сим 2.0-д шалгасан)']); s[11][0].font = Font(name=F, bold=True)
for i, st in enumerate(['Засагдсан ✓', 'Хэсэгчлэн*', 'Хойшлогдсон', 'Дахин шалгаагүй']):
    r = 12 + i; s.cell(r, 1, st); s.cell(r, 2, f'=COUNTIF(ENT!$Q:$Q,A{r})' if '*' not in st else '=COUNTIF(ENT!$Q:$Q,"Хэсэгчлэн*")')
s.append([]); s.append(['SIM2 модулиар']); s[18][0].font = Font(name=F, bold=True)
mods = sorted({i['module'].split(' (')[0].split(' / ')[0] for i in sim2})
for i, mo in enumerate(mods):
    r = 19 + i; s.cell(r, 1, mo); s.cell(r, 2, f'=COUNTIF(SIM2!$C:$C,"{mo}*")')
s.column_dimensions['A'].width = 40; s.column_dimensions['B'].width = 14; s.column_dimensions['C'].width = 14; s.column_dimensions['D'].width = 10
for row in s.iter_rows():
    for c in row:
        if not c.font.name: c.font = Font(name=F)

# ── ENT ──
e = out.create_sheet('ENT')
e.append(hdr + ['Сим 2.0 регресс', 'Регрессийн нотолгоо'])
for d in ent:
    st, ev = REG.get(d['ID'], ('Дахин шалгаагүй', ''))
    e.append([d.get(h) for h in hdr] + [st, ev])
style_header(e, len(hdr) + 2)
for row in e.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name=F); c.alignment = Alignment(wrap_text=True, vertical='top')
    if row[1].value in SEVFILL: row[1].fill = PatternFill('solid', fgColor=SEVFILL[row[1].value])
    if row[16].value and '✓' in str(row[16].value): row[16].fill = PatternFill('solid', fgColor='C6EFCE')
for i, w in enumerate([9, 6, 16, 8, 10, 48, 60, 36, 30, 48, 40, 14, 24, 10, 10, 14, 20, 40]):
    e.column_dimensions[get_column_letter(i + 1)].width = w
e.auto_filter.ref = f'A1:{get_column_letter(len(hdr) + 2)}{len(ent) + 1}'

# ── SIM2 ──
t = out.create_sheet('SIM2')
cols = ['ID', 'Зэрэг', 'Модуль', 'Суваг', 'Ангилал', 'Асуудал', 'Дэлгэрэнгүй', 'Хүлээлт', 'Бодит байдал', 'Шийдэл', 'Нотолгоо', 'Байгууллага', 'Холбоотой', 'Огноо', 'Төлөв']
t.append(cols)
for i in sim2:
    t.append([i['id'], i['sev'], i['module'], i['channel'], i['category'], i['title'], i['detail'], i['expected'], i['actual'], i['solution'], i['evidence'], i['org'], i['related'], i['ts'][:10], 'Нээлттэй'])
style_header(t, len(cols))
for row in t.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name=F); c.alignment = Alignment(wrap_text=True, vertical='top')
    if row[1].value in SEVFILL: row[1].fill = PatternFill('solid', fgColor=SEVFILL[row[1].value])
for i, w in enumerate([10, 6, 18, 10, 12, 52, 64, 40, 24, 48, 30, 8, 14, 11, 10]):
    t.column_dimensions[get_column_letter(i + 1)].width = w
t.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{len(sim2) + 1}'
out.save(OUT)
print('saved', OUT, len(ent), len(sim2))
r = subprocess.run([sys.executable, '/root/.claude/skills/synced/9977f7ca-7cff-44a0-a2e3-960262b7f870_ef9fda99-5d4a-4cc4-904d-097f30f34cb4/xlsx/scripts/recalc.py', OUT, '60'], capture_output=True, text=True)
print(r.stdout[-400:], r.stderr[-200:])
